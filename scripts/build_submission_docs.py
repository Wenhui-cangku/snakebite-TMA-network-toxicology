# -*- coding: utf-8 -*-
"""Build Cover_Letter_CBI.docx + Highlights_CBI.docx (v45, 2026-10-02).
Repo-portable version: run from anywhere; output goes to 09_manuscript/submission/.
Cover letter fits one page; highlight bullets are asserted <=85 characters (journal limit)."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = Path(__file__).resolve().parents[1] / '09_manuscript/submission'
OUT.mkdir(parents=True, exist_ok=True)

def base_doc():
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Calibri'; st.font.size = Pt(10.5)
    for sec in doc.sections:
        from docx.shared import Cm
        sec.top_margin = sec.bottom_margin = Cm(1.8)
        sec.left_margin = sec.right_margin = Cm(2.2)
    return doc

# ================= cover letter =================
doc = base_doc()
def P(text, bold=False, size=None, align=None, after=6):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = bold
    if size: r.font.size = Pt(size)
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    return p

P('Wenhui Chen, MD', bold=True, after=0)
P('Department of Emergency Medicine', after=0)
P('The Second Affiliated Hospital of Guilin Medical University', after=0)
P('Guilin, Guangxi, China', after=0)
P('chenwenhui@glmu.edu.cn  |  ORCID 0009-0001-3910-2528', after=10)
P('October 2, 2026', after=10)
P('The Editor-in-Chief', after=0)
P('Chemico-Biological Interactions (Elsevier)', after=10)
P('Dear Editor-in-Chief,', after=8)
P('On behalf of all authors, I am pleased to submit our original research article, '
  '“Network-based prioritization of venom–host interactions potentially relevant to Russell’s '
  'viper–associated thrombotic microangiopathy” by Xiuzhen Pan, Xianggui Zeng, Jiali Li, and '
  'Wenhui Chen, for consideration in Chemico-Biological Interactions.')
P('Russell’s viper envenoming can trigger thrombotic microangiopathy (TMA), yet the molecular '
  'chain from venom toxins to the TMA phenotype remains fragmented. We built a '
  'documented-analysis-plan–guided, multi-layer computational pipeline linking 63 curated '
  'Daboia russelii / D. siamensis toxins to host biology: 447 evidence-tiered toxin–target edges '
  '(15 literature-curated direct targets), six-source disease-gene integration (1,042 core genes), '
  'network topology, dual-platform enrichment, protein–protein docking on three platforms with '
  'orthogonal NMA/PRODIGY descriptors, small-molecule docking of the SVMP inhibitors '
  'batimastat/marimastat and the PLA2 inhibitor varespladib, cross-species transcriptome '
  're-analysis, network robustness testing, and AOP-framework integration.')
P('Basic findings: the coagulation/fibrinolysis axis dominates every evidence layer (direct '
  'targets, hub topology, enrichment — KEGG:04610, P = 3.9 × 10⁻¹⁸, with all eleven driver genes '
  'on the coagulation arm); endothelial/inflammatory activation forms a second, '
  'transcriptome-supported pillar; and complement involvement, a widely assumed mechanism, is '
  'not supported at the direct-target or enrichment layers and is reported honestly as a '
  'negative/secondary finding. Structural analyses converged on a factor-X-activation interface '
  'for RVV-X and a template-grade svVEGF–VEGFR2 interaction, and the hydroxamate inhibitors '
  'chelated the RVV-X catalytic zinc with correct geometry.')
P('We believe this work fits the journal’s scope: it is mechanistic toxicology of chemically '
  'characterized venom toxins, addresses toxin–protein and inhibitor–protein interactions at the '
  'molecular level, and is directly relevant to a WHO-listed neglected tropical disease. The '
  'entire pipeline is open and reproducible: code and result tables are archived on GitHub and '
  'Zenodo (DOIs in the manuscript), with a version-by-version change log and a formal '
  'Supplementary Methods and Tables S1–S10 package.')
P('The manuscript is original, has not been published previously, and is not under consideration '
  'by any other journal. All authors have read and approved the submission and declare no '
  'competing interests. No new human or animal experiments were conducted; the study reanalyzes '
  'publicly available data. Research data are deposited and linked in accordance with the '
  'journal’s research-data policy. The use of AI-assisted tools in the writing process is '
  'declared in the manuscript. We would be glad to suggest qualified reviewers through the '
  'submission system.')
P('Thank you for considering our work. We look forward to your response.', after=10)
P('Sincerely,', after=2)
P('Wenhui Chen', bold=True, after=0)
P('on behalf of all authors', after=0)
doc.save(OUT / 'Cover_Letter_CBI.docx')
print('cover letter saved')

# ================= highlights =================
hl = [
 'Multi-layer network-toxicology pipeline from Daboia venom toxins to TMA machinery',
 'Coagulation/platelet axis dominates target, network, and enrichment evidence',
 'Seven toxin–host complexes cross-checked on three docking platforms plus NMA',
 'Batimastat, marimastat and varespladib bind toxin pockets with sound geometry',
 'Cross-species transcriptomes and AOP mapping bound the mechanism honestly',
]
for h in hl:
    assert len(h) <= 85, (len(h), h)
doc = base_doc()
p = doc.add_paragraph(); r = p.add_run('Highlights'); r.bold = True; r.font.size = Pt(14)
p.paragraph_format.space_after = Pt(10)
p = doc.add_paragraph()
r = p.add_run('Network-based prioritization of venom–host interactions potentially relevant to '
              'Russell’s viper–associated thrombotic microangiopathy')
r.italic = True; r.font.size = Pt(10)
p.paragraph_format.space_after = Pt(10)
for h in hl:
    doc.add_paragraph(h, style='List Bullet')
doc.save(OUT / 'Highlights_CBI.docx')
print('highlights saved; lengths:', [len(h) for h in hl])
