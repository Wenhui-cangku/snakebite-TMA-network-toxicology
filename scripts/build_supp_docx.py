# Repo-portable copy of the workspace builder (paths resolved from this file location).
# -*- coding: utf-8 -*-
"""Build Supplementary_Material.docx (v44, 2026-10-02) — CBI submission supplement."""
from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

WS = Path(__file__).resolve().parents[1]  # repo root
REPO = Path(__file__).resolve().parents[1]  # repo root (English CSVs)
OUT = WS / '09_manuscript/supplementary/Supplementary_Material.docx'

doc = Document()
st = doc.styles['Normal']
st.font.name = 'Calibri'; st.font.size = Pt(10.5)
for i in range(1, 4):
    h = doc.styles[f'Heading {i}']
    h.font.name = 'Calibri'; h.font.color.rgb = RGBColor(0x1F, 0x3B, 0x73)

def P(text, bold=False, size=None, align=None, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    if size: r.font.size = Pt(size)
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p

def bullets(items):
    for it in items:
        doc.add_paragraph(it, style='List Bullet')

def mini_table(df, cols, widths=None, font=8.5):
    tb = doc.add_table(rows=1, cols=len(cols))
    tb.style = 'Light Grid Accent 1'
    for j, c in enumerate(cols):
        cell = tb.rows[0].cells[j]
        cell.text = str(c)
        for r in cell.paragraphs[0].runs: r.bold = True; r.font.size = Pt(font)
    for _, row in df.iterrows():
        cells = tb.add_row().cells
        for j, c in enumerate(cols):
            cells[j].text = str(row[c])
            for p in cells[j].paragraphs:
                for r in p.runs: r.font.size = Pt(font)
    return tb

# ================= cover =================
P('Supplementary Material', bold=True, size=22, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)
P('Network-based prioritization of venom–host interactions potentially relevant to '
  'Russell’s viper–associated thrombotic microangiopathy', bold=True, size=14,
  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)
P('Xiuzhen Pan¹, Xianggui Zeng¹, Jiali Li¹, Wenhui Chen¹,*', size=11,
  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
P('¹ Department of Emergency Medicine, The Second Affiliated Hospital of Guilin Medical University, Guilin, China',
  size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
P('* Corresponding author: chenwenhui@glmu.edu.cn (ORCID 0009-0001-3910-2528); '
  'first author ORCID 0009-0006-3040-7153', size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)
P('Contents: Supplementary Methods (SM.1–SM.9) · Tables S1–S10 (captions below; full machine-readable tables in '
  'the companion workbook Supplementary_Tables_S1-S10.xlsx) · companion data files note.',
  size=10, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
P('Compiled 2026-10-02 (analysis change log v44). The analysis code and result tables are archived at '
  'https://doi.org/10.5281/zenodo.23065792 (release v1.0.2); large raw files at '
  'https://doi.org/10.5281/zenodo.23087108.', size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

# ================= SM.1 =================
doc.add_heading('Supplementary Methods', level=1)
doc.add_heading('SM.1  Venom-toxin library construction (Phase 1A)', level=2)
P('Retrieval (all 2026-09-22, UniProtKB release 2026_03): main query '
  'keyword:"Toxin" AND reviewed:true AND organism_name:"Daboia russelii" returned 52 entries '
  '(28 D. siamensis covered by automatic synonym expansion; an independent D. siamensis export of 28 entries '
  'overlapped completely). Because the "Toxin" keyword misses the Kunitz inhibitor family, a supplementary query '
  '(organism_name:"Daboia russelii") AND (protein_name:"Kunitz") AND (reviewed:true) recovered 24 reviewed '
  'Kunitz entries, plus 4 CRISP entries (full-length CRISPs unreviewed, flagged in the source column). '
  'SVMP coverage was replenished from NCBI Protein with russelysin (GenBank AAZ39880) and two disintegrins. '
  'Merging and de-duplication by accession gave 78 records; 16 fragmented sequences <30 aa were removed by design '
  '(mostly short Kunitz/CRISP peptides), yielding 62 entries; the VenomZone cross-check recovered one full-length '
  'basic PLA2-3 (P86368), giving the final 63-entry library.')
P('Quality control against VenomZone/Tox-Prot (query taxonomy_id:{8707/343250} AND '
  'cc_tissue_specificity:venom AND reviewed:true; 71 reviewed venom entries): the 21 entries absent from our '
  'library were 20 short-peptide sequencing-evidence entries (<30 aa, excluded by design) plus P86368 (recovered); '
  'seven reviewed entries present in our library but lacking a venom-tissue annotation were retained as reasonable '
  'additions. All five major families (SVMP/SVSP/PLA2/snaclec/Kunitz) are represented in both resources.')
P('Redundancy was reduced by 90%-identity clustering (custom greedy CD-HIT-equivalent; script build_toxin_lib.py), '
  'giving 40 clusters with 40 representative sequences (toxins_nr90.fasta). Family-level venom abundance weights '
  'were back-filled from the functional venomics dataset of Sri Lankan D. russelii (Tan et al., J Proteomics 2015): '
  'PLA2 35.0%, snaclec 22.4%, SVSP 16.0%, SVMP 6.9%, LAAO 5.2%, Kunitz 4.6%, NGF 3.5%, CRISP 2.0%. '
  'Final composition (63 entries): Kunitz 21, PLA2 16 (including P86368 recovered at the VenomZone cross-check), '
  'snaclec/CTL 8, SVMP/disintegrin 5, SVSP 4, CRISP 4, LAAO 2, VEGF 2, NGF 1; species split D. siamensis 38 / '
  'D. russelii 25. '
  'Structure availability: 11 PDB entries, 48 AlphaFold models, 4 without structure.')

# ================= SM.2 =================
doc.add_heading('SM.2  Six-source disease-gene integration (Phase 1B)', level=2)
P('All retrievals 2026-09-22 (OMIM 2026-09-24). Phenotype scope: thrombotic microangiopathy (TMA), '
  'thrombotic thrombocytopenic purpura (TTP), atypical/typical hemolytic uremic syndrome (aHUS/HUS), '
  'microangiopathic hemolytic anemia (MAHA), snakebite envenoming, and venom-induced consumption coagulopathy '
  '(VICC). snakebite envenoming/VICC have no entries in Open Targets, DisGeNET, or OMIM (verified); '
  'the snakebite-specific layer is therefore carried by the toxin–target edge table (SM.3).')
bullets([
 'Open Targets Platform (GraphQL API): 8 MONDO terms (TMA; TTP ×3; aHUS ×3; HUS), association score ≥ 0.05 → '
 '2,014 records → 741 unique genes (backbone).',
 'DISEASES (JensenLab knowledge file, curated channels UniProtKB-KW/MedlinePlus): 15 core genes, all overlapping.',
 'GeneCards (manual browser export, six terms; per-term median score thresholds — TMA 1.39, aHUS 113.85, '
 'TTP 47.79, MAHA 107.24, snakebite 20.40, VICC 46.30): 1,356 unique genes, 698 above threshold, 1,075 newly merged.',
 'DisGeNET (academic API, CURATED sources only; 16 UMLS CUIs incl. C2717961/C0034155/C2931788/C0019061; '
 'score ≥ 0.1): 46 unique genes, all already present; 2 entries (score ≥ 0.3) promoted to the core set.',
 'CTD (batch query, manual export; disease whitelist: Snake Bites / aHUS / HUS / TTP incl. acquired / '
 'Thrombotic Microangiopathies / Anemia, Hemolytic to block CTD term expansion): direct-evidence genes → core '
 '(80 unique, 3 new); InferenceScore ≥ 50 → broad set (349 new); snakebite-envenoming term fully retained (57 genes, 13 new).',
 'OMIM (API; 44 MIM entries from five phenotype terms, phenotype whitelist of 14 entries; mappingKey = 3 only): '
 '15 genes, all already in the core set.'])
P('Nine a priori four-axis molecules not covered by any source were force-included per the documented analysis '
  'plan: F2, F5, F10, FGG, PROC, VCAM1, ICAM1, SELE, HAVCR1 (sensitivity analysis in Table S6). '
  'Core-set rule: GeneCards ≥ per-term median ∪ OpenTargets score ≥ 0.3 ∪ present in ≥ 3 disease terms ∪ '
  'DISEASES curated ∪ DisGeNET ≥ 0.3 ∪ CTD direct ∪ OMIM (mappingKey 3) ∪ a priori molecules. '
  'Final library: 2,190 genes (core 1,042 / broad 1,148); all 27/27 a priori molecules are in the core set. '
  'Symbols were mapped to human UniProt AC with mygene (species=human).')

# ================= SM.3 =================
doc.add_heading('SM.3  Toxin–host edge mapping and four-tier evidence (Phase 2)', level=2)
P('Edges were built in four evidence tiers. L1 (curated interaction databases CTD/T3DB/BindingDB): 0 edges — '
  'CTD blocked programmatic access (ALTCHA) and no equivalent curated toxin–target rows exist for these toxins. '
  'L2 (literature-curated): 49 edges / 15 unique human targets, extracted from UniProt/Swiss-Prot curated '
  'FUNCTION annotations of the 63 entries (43 entries with target-related descriptions) plus targeted PubMed '
  'verification (PMIDs 27089306, 21356226, 18554518, 28042812, 36423674, 37092784, 28732041, 21871889, 18222185, '
  '21802487); family-level edges are flagged (score 0.6) when a publication could not be assigned to a specific '
  'library entry. L3 (BLAST homology auxiliary): suspended (no local blastp; plan-labelled auxiliary only). '
  'L4 (STRING one-hop expansion): for the 15 direct targets, STRING API interaction_partners (H. sapiens, '
  'combined score ≥ 0.7, top 10) gave 105 unique neighbours → 398 edges. Total 447 edges / 112 unique targets. '
  'Entries without curated human targets are recorded as no-edge with reasons (CRISP ×4, LAAO ×2, chaperone PLA2 '
  'P31100, snaclec Q4PRC6–9/Q4PRD0).')

# ================= SM.4 =================
doc.add_heading('SM.4  Network construction, topology, and robustness (Phase 3)', level=2)
P('Intersection: 15 direct targets (L1+L2) ∩ 1,042 core disease genes → strict core of 10 (COL4A1, F10, F11, F5, '
  'FGA, GP1BA, KDR, PLG, PROC, PROS1). An extended sensitivity/topology layer was built as (strict core ∪ '
  'L4 STRING neighbours ∩ core set) = 54 nodes. PPI edges were exported from the STRING API '
  '(string-db.org/api, STRING v12.0, queried 2026-09-24) at combined score ≥ 0.7 (main; 321 edges) and ≥ 0.4 '
  '(sensitivity; 615 edges). Hub ranking used a Python re-implementation of cytoHubba Degree, MCC (maximal clique '
  'centrality) and EPC (edge percolated component; p = 0.5, 1,000 simulations, seed 42), Top-20 consensus → 10 hub '
  'genes. EPC had low discrimination on this dense coagulation network (scores converge; stated in Methods), so the '
  'consensus is driven by MCC/Degree. Robustness: targeted removal in static degree order (computed once on the '
  'full network) versus 100 random-removal shuffles (seed 1234); curves include k = 0 (LCC/N = 1.0) as the first '
  'point; at 27/54 nodes removed (50%), LCC = 0.241 (targeted) vs 0.480 ± 0.023 (random).')

# ================= SM.5 =================
doc.add_heading('SM.5  Dual-platform enrichment (Phase 4)', level=2)
P('Input: Hub 10 ∪ strict core 10 = 15 unique genes. Platform 1: g:Profiler g:GOSt (API, 2026-09-24; '
  'organism hsapiens; sources GO:BP/CC/MF + KEGG + Reactome; g:SCS correction, user threshold 0.05; no custom '
  'background) → 116 significant terms. Platform 2: Enrichr (KEGG_2021_Human, Reactome_2022, '
  'GO_Biological_Process_2023; BH correction) → 202 significant terms. Cross-platform correspondences are '
  'theme-level (not identical term IDs); nine themes matched by keyword. A re-run of the 15-gene g:Profiler query '
  'on 2026-10-02 reproduced the archived 116-term set exactly (reproducibility check, Table S6).')

# ================= SM.6 =================
doc.add_heading('SM.6  Structure validation and small-molecule docking (Phase 5)', level=2)
P('Structures (Table S7): experimental PDB entries 2E3X (RVV-X heavy chain), 3S9C (RVV-Vγ–FV complex), 1KPM '
  '(PLA2), 1WQ9 (svVEGF), 2W26 (FXa), 7KVE (FV), 3GHG (fibrinogen fragment), 3V2A (VEGFR2), 1QRZ (plasmin), '
  '1PPB; AlphaFold models (pLDDT > 70 rule) for daborhagin-K, Kunitz inhibitor, snaclec, and 1M10 was used only '
  'before retraction (see below). ClusPro (Balanced mode, blind docking): pre-specified acceptance = top-1 cluster '
  '≥ 30 members; 7/7 pairs passed (59–190). HADDOCK 2.4 (biochemical AIR restraints; interface cutoff 5 Å): '
  'pre-specified double pass = score ≤ −100 and functional-site coverage; P2 passed both (score −144.9 ± 4.4, '
  'Z = −1.4, 6/6); P1/P3/P6 achieved restraint orientation (nearest restraint distance 1.8–7.8 Å) without '
  'reaching the score threshold and are reported as site-orientation evidence only. HDock (template docking, P7): '
  'confidence 0.8969 (> 0.7), 21/26 reference-interface coverage, receptor superposition on 3V2A RMSD 0.000 Å. '
  'Orthogonal descriptors (post hoc, labelled): iMODS NMA mobility ratios and PRODIGY ΔG/Kd for the 7 pairs.')
P('Retraction note: pair P4 (snaclec × "GP1BA") was withdrawn on 2026-10-02 after the receptor chain was found '
  'misidentified — 1M10 chain A is the VWF-A1 domain, not GP1BA; the pair is removed from the evidence chain '
  '(manuscript protocol-deviation table; Fig. 6 panel D blanked).')
P('Small molecules (Table S8): AutoDock Vina 1.2.7 + Meeko 0.8.0; ligands from PubChem 3D (batimastat CID '
  '5362422, marimastat CID 119031, varespladib CID 155815); receptors RVVX_A (2E3X chain A, catalytic Zn '
  'retained, +2), daborhaginK (AlphaFold + modelled Zn at His339/343/349 — arm not trusted), PLA2_1KPM (chain A, '
  'grid centred on His48/Asp49); grid box 26 × 26 × 26 Å (centres in Table S8), exhaustiveness = 8, num_modes = 9, '
  'default seed. Hydroxamate Zn-chelation geometry was verified on the top poses (Table S8).')

# ================= SM.7 =================
doc.add_heading('SM.7  GEO transcriptome re-analysis (v41 canonical pipeline)', level=2)
P('GSE121297 (human HUVEC × rhodocetin-αβ; Affymetrix HG-U133 Plus 2.0 / GPL570): static RCab (GSM3430997–99) '
  'vs static control (GSM3430994–96), 3 vs 3; the 80-rpm arm was not used. GSE248215 (mouse skeletal muscle, '
  'D. russelii venom 30 µg in vivo; NanoString Fibrosis V2, 760 genes): DR 24 h (3) vs PBS control (2); '
  'author-normalised log2 matrix used. Differential expression: Welch’s t-tests in Python with '
  'Benjamini–Hochberg correction (script geo_v41_reanalysis.py). Probe handling: one fixed probe per gene, '
  'highest mean across the six contrast samples chosen before testing; a probe mapping to several symbols '
  'contributes to each symbol → 22,453 genes. An earlier archived run (21,183 genes; multi-symbol probe conflicts '
  'resolved differently) gives identical log2FC/raw P for all shared genes and the same conclusion: 0 genes at '
  'BH-adjusted P < 0.05 in either dataset. All interpretation is nominal, exploratory, and labelled as such; the '
  'two tissues are not combined into a single validation cohort.')

# ================= SM.8 =================
doc.add_heading('SM.8  Protocol deviations (mirrors manuscript Table 1)', level=2)
dev = pd.DataFrame([
 ['Hypotheses','H1 SVMP–VWF/ADAMTS13 dominant; H2 complement & endothelial enrichment; exploratory H3','H1 coagulation/platelet prominence; H2 complement as secondary host response; H3 dropped','Intersection and enrichment results motivated reformulation','Yes — reformulated hypotheses are post hoc and labelled'],
 ['Module screening','MCODE cluster screening','MCC/Degree/EPC Top-20 consensus (re-implementation)','MCODE step not executed in the final pipeline','Yes'],
 ['Docking controls','3–5 negative docking pairs','Not performed','Resource constraints; listed as a limitation','n/a'],
 ['Enrichment platforms','Metascape + clusterProfiler','g:Profiler + Enrichr','Web-platform accessibility and full-request auditability','Yes'],
 ['GEO statistics','limma (R)','Welch’s t-tests + BH in Python','Python-only analysis environment; method renamed accordingly','Yes'],
 ['Orthogonal corroboration','Not planned','PRODIGY + iMODS model-property descriptors','Added during revision; reported as post hoc descriptors','Yes (post hoc addition)'],
], columns=['Item','Original documented plan','Final analysis','Reason for change','Results seen before change?'])
mini_table(dev, list(dev.columns), font=7.5)

# ================= SM.9 =================
doc.add_heading('SM.9  Reproducibility manifest', level=2)
bullets([
 'Environment: Python 3.12 (pandas, numpy, scipy, matplotlib, seaborn, networkx, openpyxl); web services — '
 'ClusPro, HADDOCK 2.4, HDock, iMODS, PRODIGY, g:Profiler, Enrichr, STRING v12.0; AutoDock Vina 1.2.7 + Meeko '
 '0.8.0; PyMOL (figures); Cytoscape 3.x (Fig. 5 rendering).',
 'Fixed random seeds: 1234 (robustness shuffles), 42 (EPC percolation); all pipeline scripts under scripts/ and '
 'phase directories; reproduction order (13 steps) in the repository README.',
 'Archive: GitHub https://github.com/Wenhui-cangku/snakebite-TMA-network-toxicology · Zenodo concept DOI '
 '10.5281/zenodo.23065792 (release v1.0.2 = 10.5281/zenodo.23097132) · large raw files dataset '
 '10.5281/zenodo.23087108.',
 'Change log: version-by-version record v1–v44 (CHANGELOG.csv in the repository; public copy of the data-'
 'management table in 00_protocol/).',
 'Reproducibility checks passed: 15-gene g:Profiler rerun = archived 116 terms (2026-10-02); robustness script '
 'rerun reproduces archived curves value-by-value; GEO script reproduces the canonical 22,453-gene table.'])

# ================= Table captions =================
doc.add_page_break()
doc.add_heading('Supplementary Tables S1–S10', level=1)
P('Full machine-readable tables are provided in the companion workbook Supplementary_Tables_S1-S10.xlsx '
  '(one sheet per table; sheet names in parentheses). Key tables are also reproduced inline below when compact.',
  space_after=10)

def caption(num, title, body):
    p = doc.add_paragraph()
    r = p.add_run(f'Table {num}. {title}. ')
    r.bold = True
    doc.paragraphs[-1].add_run(body)

caption('S1','Venom-toxin master table (sheet S1_Toxin_master; 63 entries)',
 'entry_id, toxin_name, toxin_family, UniProt/GenBank ID, entry_name, species, sequence length, source '
 '(incl. unreviewed/NCBI flags), cluster ID (40 clusters, 90% identity), structure status (PDB/AlphaFold/none), '
 'family-level venom abundance weight (%), inclusion status and exclusion reason, download date. Companion FASTA '
 'files (all_toxins_raw.fasta, toxins_nr90.fasta) are in the Zenodo raw dataset. Cited in Methods §2.1 and Fig. 1.')
caption('S2','Integrated disease-gene library (sheet S2_Disease_genes; 2,190 records)',
 'gene_symbol, UniProt AC, Entrez ID, per-source evidence (GeneCards score, DisGeNET score, OMIM yes/no, '
 'CTD inference, OpenTargets score, OT disease count, DISEASES curated), core-set flag (1,042 core / 1,148 broad), '
 'pathology-axis assignment (1–4), source query terms, download date. Cited in Methods §2.2 and Fig. 2.')
caption('S3','Toxin–host edge table with four-tier evidence (sheet S3_Toxin_target_edges; 447 edges)',
 'toxin_id, toxin_ac, family, target_symbol, evidence_type (L2 literature-curated / L4 STRING neighbour), '
 'evidence_detail (effect direction, isoform/complex notes, experimental context, or “via <target>” for L4), '
 'PMIDs, evidence score, download date. L1 = 0 and L3 suspended (SM.3). Cited in Methods §2.3 and Figs. 2–5.')
caption('S4','Full three-algorithm hub ranking (sheet S4_Hub_full_ranking; 54 nodes)',
 'Degree, MCC, and EPC scores with competition ranks (ties share the minimum rank, reported in full) for every '
 'node of the extended layer, plus hub-10 / strict-core-10 flags and pathology-axis labels. Network settings in '
 'SM.4; EPC low-discrimination note applies. Cited in Results (hub analysis) and Table 2.')
s4 = pd.read_csv(REPO / '04_network/hub_genes.csv', encoding='utf-8-sig').fillna('—')
P('Hub-10 excerpt (full 54-node ranking in the workbook):', space_after=2)
mini_table(s4, list(s4.columns), font=8)
doc.add_paragraph()
caption('S5','Enrichment results (sheets S5a_gProfiler, S5b_Enrichr, S5c_Prior_pathways)',
 'S5a: all 116 g:Profiler terms significant at g:SCS < 0.05 (15-gene input) with term IDs, sizes and driver genes. '
 'S5b: all 424 Enrichr rows across the three libraries, including non-significant terms (rank, P, BH q, overlap, '
 'genes). S5c: the seven a priori pathway tests on both platforms with conclusions and driver genes. '
 'Cited in Results §3.4 and Fig. 4 / Table 3.')
s5c = pd.read_csv(REPO / '05_enrichment/prior_pathway_check.csv', encoding='utf-8-sig').fillna('—')
mini_table(s5c, list(s5c.columns), font=8)
doc.add_paragraph()
caption('S6','Sensitivity analysis without forced inclusion (sheets S6_Sensitivity, S6_gProfiler_11genes)',
 'Documents the effect of removing the nine force-included a priori genes: input sets, disease-set sizes '
 '(1,042 → 1,033 core), strict-core/hub composition, and a full g:Profiler rerun on the 11-gene input '
 '(93 terms). The dominant coagulation/platelet theme is qualitatively stable (fibrin clot 7/11, P = 3.8e-13; '
 'KEGG:04610 7/11, P = 1.4e-10); gamma-carboxylation becomes marginal. The sheet also records the 15-gene '
 'reproducibility rerun (116/116 terms identical). Cited in Results §3.2.')
caption('S7','Structure-validation summary (sheets S7_Docking, S7_iMODS_PRODIGY)',
 'Seven toxin–host pairs across ClusPro (blind, Balanced), HADDOCK 2.4 (AIR-restrained) and HDock (template): '
 'cluster sizes, energies/scores, Z-scores, interface residue counts, functional-site coverage, restraint '
 'distances, and verdicts, plus the iMODS/PRODIGY orthogonal descriptors. The retracted P4 pair is listed as a '
 'marked retraction row. AIR restraint files, PDB/mmCIF inputs and output models are in the Zenodo raw dataset. '
 'Cited in Results §3.5–3.6, Figs. 6–7, Tables 4–5.')
dock = pd.read_csv(REPO / '07_docking/docking_results.csv', encoding='utf-8-sig', engine='python', on_bad_lines='skip')
P('Compact view of the iMODS/PRODIGY orthogonal descriptors (full ClusPro/HADDOCK/HDock metrics and notes in the workbook):', space_after=2)
clu = pd.read_csv(REPO / '08_md/imods_prodigy_summary.csv', encoding='utf-8-sig')
mini_table(clu, list(clu.columns), font=8)
doc.add_paragraph()
caption('S8','Small-molecule docking (sheet S8_Small_molecule)',
 'Nine receptor × ligand results (batimastat/marimastat/varespladib × RVV-X/daborhagin-K/PLA2): best Vina '
 'affinities, count of modes ≤ −7.5 kcal/mol, grid centres (26 Å box), and the Zn-chelation geometry check; '
 'the daborhagin-K AlphaFold + modelled-Zn arm is marked not trusted. Settings in SM.6. Cited in Results §3.7 '
 'and Fig. 8.')
caption('S9','GEO re-analysis (sheets S9a_GEO_samples, S9b_GSE121297_DEG, S9c_GSE248215_DEG, S9d_GEO_watchlist)',
 'S9a: sample tables and contrast design (GSM IDs, titles, platforms, groups used). S9b/S9c: full per-gene '
 'Welch + BH tables (22,453 and 760 genes; v41 canonical fixed-probe pipeline). S9d: the 35-molecule a priori '
 'watchlist across both datasets (log2FC, P, significance flags; ALB and COL4A1 not measured by either platform). '
 'Cited in Results §3.8, Figs. 9–10, Table 6.')
caption('S10','Network robustness (sheets S10_Robustness, S10_Removal_order)',
 'k = 0…54 removal axis (k = 0 included); the static-degree targeted curve, all 100 per-run random-removal curves '
 '(seed 1234), and mean ± SD columns; the targeted removal order with degrees. Regeneration script: '
 '04_network/robustness_analysis_v41.py. Cited in Results §3.9 and Fig. 11.')

# ================= companion files =================
doc.add_heading('Companion data files', level=1)
bullets([
 'Zenodo raw dataset (10.5281/zenodo.23087108): toxin FASTA files; PDB/mmCIF docking inputs and output models '
 '(ClusPro/HADDOCK/HDock); GEO raw matrices and GPL570 probe annotation; iMODS/PRODIGY raw outputs; CTD raw TSV '
 'exports; file inventory in README_DATA.md inside the zip.',
 'Code + result tables: GitHub repository and Zenodo software archive (concept DOI 10.5281/zenodo.23065792; '
 'release v1.0.2 = 10.5281/zenodo.23097132), including the change log (v1–v44) and the public data-management table.',
 'Cytoscape session files for Fig. 5 (cyjs/xgmml) and the vector SVG export: repository 04_network/cytoscape/.',
 'Operational SOP for future MD/free-energy work (planned, not executed): repository 08_md/MD_SOP.md — '
 'explicitly a future-work protocol, not completed evidence.'])

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print('saved:', OUT)
