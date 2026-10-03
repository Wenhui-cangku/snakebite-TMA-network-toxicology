# -*- coding: utf-8 -*-
"""v53: reference DOIs, 2 new refs, citation insertions, HAVCR1 fix in 4.5."""
import shutil
from docx import Document

SRC = '蛇伤TMA论文_CBI投稿版_EN_v41_revised.docx'
shutil.copy(SRC, SRC + '.bak_v53')
d = Document(SRC)
ps = d.paragraphs

def set_text(p, new):
    runs = p.runs
    assert runs, 'no runs'
    runs[0].text = new
    for r in runs[1:]:
        r.text = ''

def replace_in(idx, old, new):
    t = ps[idx].text
    assert old in t, 'NOT FOUND [%d]: %s' % (idx, old[:60])
    set_text(ps[idx], t.replace(old, new, 1))

# --- DOI additions (ref number -> paragraph index 125+n-1) ---
DOIS = {
 11:'10.1002/etc.60', 12:'10.1093/nar/gkae1010', 15:'10.1002/cpbi.5',
 16:'10.1093/nar/gkaa1027', 17:'10.1016/j.ymeth.2014.11.020', 18:'10.1093/nar/gkz1021',
 19:'10.1093/nar/gkaa891', 20:'10.1093/nar/gky1151', 21:'10.1186/s13059-016-0953-9',
 23:'10.1186/1752-0509-8-S4-S11', 25:'10.1093/nar/gkz369', 26:'10.1186/1471-2105-14-128',
 27:'10.1093/nar/gkw377', 28:'10.1038/75556', 29:'10.1093/nar/28.1.27',
 30:'10.1093/nar/gkab1028', 31:'10.1111/j.2517-6161.1995.tb02031.x', 32:'10.1093/nar/28.1.235',
 33:'10.1038/s41586-021-03819-2', 34:'10.1038/nprot.2016.169', 35:'10.1016/j.jmb.2015.09.014',
 37:'10.1038/s41596-020-0312-x', 38:'10.1093/nar/gkx407', 39:'10.1002/jcc.21334',
 40:'10.1021/acs.jcim.1c00203', 41:'10.1093/bioinformatics/btw514', 42:'10.1093/nar/gku339',
 43:'10.1093/nar/30.1.207', 46:'10.1101/gr.1239303',
}
for n, doi in DOIS.items():
    p = ps[125 + n - 1]
    t = p.text
    assert t.startswith('[%d]' % n), 'ref mismatch at %d' % n
    assert 'doi.org' not in t, 'already has doi: %d' % n
    set_text(p, t.rstrip() + ' https://doi.org/%s.' % doi)

# --- [22] STRING updated to the 2025 paper (matches v12.5 in Methods) ---
replace_in(146,
    "D. Szklarczyk, A.L. Gable, K.C. Nastou, D. Lyon, R. Kirsch, S. Pyysalo, N.T. Doncheva, M. Legeay, T. Fang, P. Bork, L.J. Jensen, C. von Mering, The STRING database in 2023: protein\u2013protein association networks and functional enrichment analyses for any sequenced genome of interest, Nucleic Acids Res. 51 (D1) (2023) D638\u2013D646.",
    "D. Szklarczyk, R. Kirsch, M. Koutrouli, K. Nastou, F. Mehryary, R. Hachilif, A.L. Gable, T. Fang, N.T. Doncheva, S. Pyysalo, P. Bork, L.J. Jensen, C. von Mering, The STRING database in 2025: protein networks with directionality of regulation, Nucleic Acids Res. 53 (D1) (2025) D730\u2013D737. https://doi.org/10.1093/nar/gkae1113.")

# --- new refs [53] [54] appended after [52] ---
ref_style = ps[176].style
p53 = d.add_paragraph("[53] G.V. Rudresha, S. Khochare, N.R. Casewell, K. Sunagar, Preclinical evaluation of small molecule inhibitors as early intervention therapeutics against Russell's viper envenoming in India, Commun. Med. 5 (1) (2025) 226. https://doi.org/10.1038/s43856-025-00900-z.", style=ref_style)
p54 = d.add_paragraph("[54] P. Ojuka, G.S. Nyamato, C.B.R. Santos, N.M. Kimani, A review and in silico screening of plant-derived snake venom/toxin inhibitors: ADMET, drug-likeness, and medicinal chemistry profiling, PLoS Negl. Trop. Dis. 19 (10) (2025) e0013579. https://doi.org/10.1371/journal.pntd.0013579.", style=ref_style)

# --- citation insertions ---
replace_in(10,
    "offers one way to map this space and to make the untested links explicit.",
    "offers one way to map this space and to make the untested links explicit. Computational prioritization around venom toxins has recent precedents in in silico inhibitor screening studies [54].")
replace_in(103,
    "hydroxamate broad-spectrum metalloproteinase inhibitors with existing preclinical/clinical data",
    "hydroxamate broad-spectrum metalloproteinase inhibitors with existing preclinical/clinical data [53]")
replace_in(103,
    "HAVCR1 showed a positive estimated expression change (log2FC = 0.99; nominal P = 0.144; BH-adjusted P = 0.378).",
    "HAVCR1 showed a positive estimated expression change (descriptive log2FC = 0.99; no statistical testing, Section 2.7).")

d.save(SRC)
print('v53 refs saved OK; total paragraphs now', len(d.paragraphs))
